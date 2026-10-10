"""Erzeugt Kahoot-Importdateien (Excel, Aufbau wie die offizielle Kahoot-Vorlage) aus inhalte/vokabeln.tsv.
Pro Buchabschnitt je eine Datei Englisch → Deutsch und Deutsch → Englisch.
Aufruf: python3 werkzeuge/kahoot_export.py   ->  export/kahoot/*.xlsx
"""
import csv, os, random, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

BASIS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ZIEL = os.path.join(BASIS, "export", "kahoot")
ZEIT = 20  # Sekunden pro Frage (Kahoot erlaubt 5, 10, 20, 30, 60, 90, 120, 240)
ARIAL, FETT = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)


def kurz(text, max_len):
    return text if len(text) <= max_len else text[: max_len - 1] + "…"


def datei(name, fragen):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kahoot"
    ws["B2"] = "Quiz template"
    ws["B2"].font = Font(name="Arial", size=14, bold=True)
    ws["B3"] = f"{name} – erzeugt aus inhalte/vokabeln.tsv. Import in Kahoot: Neues Kahoot → Fragen aus Tabelle importieren."
    ws["B4"] = "Hinweis: Aufbau wie die Kahoot-Vorlage (Daten ab Zeile 9). Meldet Kahoot einen Fehler, die Zeilen 9 ff. in die offizielle Vorlage kopieren."
    for z in ("B3", "B4"):
        ws[z].font = ARIAL
    kopf = ["", "Question - max 120 characters", "Answer 1 - max 75 characters", "Answer 2 - max 75 characters",
            "Answer 3 - max 75 characters", "Answer 4 - max 75 characters",
            "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs", "Correct answer(s) - choose at least one"]
    for s, t in enumerate(kopf, 1):
        c = ws.cell(row=8, column=s, value=t)
        c.font, c.fill = FETT, PatternFill("solid", fgColor="E8E0F5")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, (frage, antworten, richtig) in enumerate(fragen, 1):
        werte = [i, frage, *antworten, ZEIT, richtig]
        for s, w in enumerate(werte, 1):
            ws.cell(row=8 + i, column=s, value=w).font = ARIAL
    for s, b in zip("ABCDEFGH", (5, 48, 30, 30, 30, 30, 14, 14)):
        ws.column_dimensions[s].width = b
    ws.row_dimensions[8].height = 45
    wb.save(os.path.join(ZIEL, name + ".xlsx"))


def main():
    random.seed(3)  # gleiche Dateien bei jedem Lauf
    os.makedirs(ZIEL, exist_ok=True)
    with open(os.path.join(BASIS, "inhalte", "vokabeln.tsv"), encoding="utf-8") as f:
        alle = list(csv.DictReader(f, delimiter="\t"))
    abschnitte = []
    for v in alle:
        if not abschnitte or abschnitte[-1][0] != v["abschnitt"]:
            abschnitte.append((v["abschnitt"], []))
        abschnitte[-1][1].append(v)
    for nr, (abschnitt, vok) in enumerate(abschnitte, 1):
        for richtung in ("EN-DE", "DE-EN"):
            frage_feld, antwort_feld = ("englisch", "deutsch") if richtung == "EN-DE" else ("deutsch", "englisch")
            sprache = "Deutsch" if richtung == "EN-DE" else "Englisch"
            fragen = []
            for v in vok:
                # Ablenker bevorzugt aus demselben Abschnitt, sonst aus der ganzen Unit
                pool = [x[antwort_feld] for x in vok if x[antwort_feld] != v[antwort_feld]]
                if len(set(pool)) < 3:
                    pool = [x[antwort_feld] for x in alle if x[antwort_feld] != v[antwort_feld]]
                falsch = random.sample(sorted(set(pool)), 3)
                antworten = [kurz(a, 75) for a in falsch + [v[antwort_feld]]]
                random.shuffle(antworten)
                richtig = antworten.index(kurz(v[antwort_feld], 75)) + 1
                frage = kurz(f"Was heißt „{v[frage_feld]}“ auf {sprache}?", 120)
                fragen.append((frage, antworten, richtig))
            titel = "Start (My day)" if abschnitt.startswith("Unit 3") else abschnitt
            name = f"{nr:02d} Unit3 {titel} {richtung}"
            datei(name, fragen)
            print(name, len(fragen), "Fragen")


if __name__ == "__main__":
    main()
